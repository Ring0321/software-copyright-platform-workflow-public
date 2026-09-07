from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from hashlib import sha256
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn


CODE_LINE = re.compile(r"^\s*\d+\s{2}")


def emu(value: Any) -> int | None:
    return int(value) if value is not None else None


def mm(value: Any) -> float | None:
    return round(float(value.mm), 3) if value is not None else None


def pt(value: Any) -> float | None:
    return round(float(value.pt), 3) if value is not None else None


def enum_value(value: Any) -> str | int | None:
    if value is None:
        return None
    return getattr(value, "name", None) or int(value)


def color_value(font: Any) -> str | None:
    try:
        rgb = font.color.rgb
        return str(rgb) if rgb is not None else None
    except (AttributeError, ValueError):
        return None


def style_profile(style: Any) -> dict[str, Any]:
    font = style.font
    paragraph = style.paragraph_format
    line_spacing = paragraph.line_spacing
    if hasattr(line_spacing, "pt"):
        line_spacing_value: float | str | None = pt(line_spacing)
    elif line_spacing is None:
        line_spacing_value = None
    else:
        line_spacing_value = str(line_spacing)
    return {
        "style_id": style.style_id,
        "name": style.name,
        "type": enum_value(style.type),
        "base_style": style.base_style.name if style.base_style is not None else None,
        "font": {
            "name": font.name,
            "size_pt": pt(font.size),
            "bold": font.bold,
            "italic": font.italic,
            "underline": enum_value(font.underline),
            "color": color_value(font),
        },
        "paragraph": {
            "alignment": enum_value(paragraph.alignment),
            "left_indent_mm": mm(paragraph.left_indent),
            "right_indent_mm": mm(paragraph.right_indent),
            "first_line_indent_mm": mm(paragraph.first_line_indent),
            "space_before_pt": pt(paragraph.space_before),
            "space_after_pt": pt(paragraph.space_after),
            "line_spacing": line_spacing_value,
            "keep_with_next": paragraph.keep_with_next,
            "page_break_before": paragraph.page_break_before,
        },
    }


def section_profile(section: Any, index: int) -> dict[str, Any]:
    return {
        "index": index,
        "orientation": "landscape" if section.orientation == WD_ORIENT.LANDSCAPE else "portrait",
        "start_type": enum_value(section.start_type),
        "page_width_emu": emu(section.page_width),
        "page_height_emu": emu(section.page_height),
        "page_width_mm": mm(section.page_width),
        "page_height_mm": mm(section.page_height),
        "top_margin_mm": mm(section.top_margin),
        "bottom_margin_mm": mm(section.bottom_margin),
        "left_margin_mm": mm(section.left_margin),
        "right_margin_mm": mm(section.right_margin),
        "header_distance_mm": mm(section.header_distance),
        "footer_distance_mm": mm(section.footer_distance),
        "header_text": [paragraph.text for paragraph in section.header.paragraphs],
        "footer_text": [paragraph.text for paragraph in section.footer.paragraphs],
        "header_linked_to_previous": section.header.is_linked_to_previous,
        "footer_linked_to_previous": section.footer.is_linked_to_previous,
    }


def table_profile(table: Any, index: int) -> dict[str, Any]:
    column_widths: list[int | None] = []
    if table.rows:
        column_widths = [emu(cell.width) for cell in table.rows[0].cells]
    shading: list[str | None] = []
    if table.rows:
        for cell in table.rows[0].cells:
            shd = cell._tc.tcPr.find(qn("w:shd"))
            shading.append(shd.get(qn("w:fill")) if shd is not None else None)
    return {
        "index": index,
        "style": table.style.name if table.style is not None else None,
        "rows": len(table.rows),
        "columns": len(table.columns),
        "first_row_column_widths_emu": column_widths,
        "first_row_shading": shading,
        "autofit": table.autofit,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract deterministic formatting from a DOCX template")
    parser.add_argument("docx", type=Path)
    parser.add_argument("output_json", type=Path)
    args = parser.parse_args()

    path = args.docx.resolve()
    document = Document(str(path))
    style_usage = Counter(paragraph.style.name for paragraph in document.paragraphs)
    used_style_names = set(style_usage)
    used_style_names.update(table.style.name for table in document.tables if table.style is not None)
    used_styles = [style for style in document.styles if style.name in used_style_names]

    page_breaks = 0
    for paragraph in document.paragraphs:
        page_breaks += len(paragraph._p.xpath('.//w:br[@w:type="page"]'))

    code_paragraphs = [paragraph for paragraph in document.paragraphs if CODE_LINE.match(paragraph.text)]
    code_samples = []
    for paragraph in code_paragraphs[:5]:
        run = paragraph.runs[0] if paragraph.runs else None
        code_samples.append(
            {
                "text": paragraph.text[:160],
                "style": paragraph.style.name,
                "alignment": enum_value(paragraph.alignment),
                "font_name": run.font.name if run is not None else None,
                "font_size_pt": pt(run.font.size) if run is not None else None,
                "font_color": color_value(run.font) if run is not None else None,
            }
        )

    profile = {
        "source": str(path),
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path.read_bytes()).hexdigest(),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
        "inline_image_count": len(document.inline_shapes),
        "page_break_count": page_breaks,
        "core_properties": {
            "title": document.core_properties.title,
            "subject": document.core_properties.subject,
            "author": document.core_properties.author,
            "category": document.core_properties.category,
        },
        "sections": [section_profile(section, index) for index, section in enumerate(document.sections)],
        "style_usage": dict(style_usage.most_common()),
        "used_styles": [style_profile(style) for style in used_styles],
        "tables": [table_profile(table, index) for index, table in enumerate(document.tables)],
        "images": [
            {
                "index": index,
                "type": enum_value(shape.type),
                "width_emu": emu(shape.width),
                "height_emu": emu(shape.height),
                "width_mm": mm(shape.width),
                "height_mm": mm(shape.height),
            }
            for index, shape in enumerate(document.inline_shapes)
        ],
        "code_format": {
            "numbered_code_paragraph_count": len(code_paragraphs),
            "samples": code_samples,
        },
    }

    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(profile, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
